"""ACDP GUI Dashboard — launched by ``python -m acdp``.

Redesigned professional dark-theme dashboard:
  • Centered ACDP hero header with subtitle
  • Version badge top-right
  • Proper color palette with contrast hierarchy
  • Agent + Connector cards with animated indicators
  • Knowledge Ingest panel
  • Resizable/draggable Live Audit Log pane
"""

from __future__ import annotations

import queue
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from typing import Callable, Any
import datetime

__all__ = ["launch_dashboard"]

# ---------------------------------------------------------------------------
# Professional colour palette
# ---------------------------------------------------------------------------
# Backgrounds
BG_BASE    = "#0a0e1a"   # deepest background
BG_SURFACE = "#111827"   # card / panel surface
BG_RAISED  = "#1a2236"   # slightly lifted element
BG_INPUT   = "#1e2d40"   # input fields
BG_HEADER  = "#0d1424"   # top header bar

# Borders
BD_SUBTLE  = "#1f2d45"   # subtle border
BD_ACCENT  = "#2d4a6e"   # highlighted border

# Text
TX_PRIMARY = "#e2e8f0"   # primary text
TX_SECONDARY = "#94a3b8" # secondary / dim text
TX_MUTED   = "#4a5568"   # very muted

# Accent colours (purposeful, not decorative)
CLR_BLUE   = "#3b82f6"   # primary action blue
CLR_GREEN  = "#10b981"   # success / active
CLR_AMBER  = "#f59e0b"   # warning / pending
CLR_RED    = "#ef4444"   # error / danger
CLR_PURPLE = "#8b5cf6"   # agents accent
CLR_CYAN   = "#06b6d4"   # connectors accent
CLR_WHITE  = "#f8fafc"   # headings

# Agent accent colours
AGENT_COLORS = {
    "Guardrail":  "#a78bfa",   # soft purple
    "Blue Team":  "#60a5fa",   # sky blue
    "Red Team":   "#f87171",   # soft red
    "DevSecOps":  "#34d399",   # emerald
}

FONT_HERO   = ("Helvetica", 36, "bold")
FONT_SUB    = ("Helvetica", 13)
FONT_H1     = ("Helvetica", 15, "bold")
FONT_H2     = ("Helvetica", 13, "bold")
FONT_BODY   = ("Helvetica", 12)
FONT_SMALL  = ("Helvetica", 11)
FONT_MONO   = ("Courier New", 11)
FONT_VER    = ("Helvetica", 11)

SPINNER = ["⣾","⣽","⣻","⢿","⡿","⣟","⣯","⣷"]
SCAN    = ["▱▱▱▱▱","▰▱▱▱▱","▰▰▱▱▱","▰▰▰▱▱",
           "▰▰▰▰▱","▰▰▰▰▰","▱▰▰▰▰","▱▱▰▰▰","▱▱▱▰▰","▱▱▱▱▰"]

# Greyed-out state colours (used when a button is inactive)
BTN_GREY_BG  = "#2d3748"
BTN_GREY_FG  = "#6b7280"
BTN_START_BG = "#16a34a"
BTN_STOP_BG  = "#dc2626"
BTN_BLUE_BG  = "#2563eb"
BTN_SLATE_BG = "#374151"


# ---------------------------------------------------------------------------
# StyledButton — works around macOS ignoring bg on disabled tk.Button
# ---------------------------------------------------------------------------
class StyledButton(tk.Label):
    """A clickable Label that looks like a button and never loses its color.

    macOS tkinter overrides tk.Button background when state='disabled'.
    By using a Label we keep full control of colors at all times.
    Visually greys itself out when deactivated; fully colored when active.
    """

    def __init__(self, parent, text: str, command: Callable,
                 active_bg: str, active_fg: str = "#ffffff",
                 font=None, **kw):
        self._cmd = command
        self._active_bg = active_bg
        self._active_fg = active_fg
        self._enabled = True
        super().__init__(
            parent, text=text,
            bg=active_bg, fg=active_fg,
            font=font or FONT_H2,
            cursor="arrow",
            relief="groove", bd=2,
            padx=14, pady=7,
            **kw,
        )
        self.bind("<Button-1>", self._on_click)
        self.bind("<Enter>",    self._on_enter)
        self.bind("<Leave>",    self._on_leave)

    def _on_click(self, _e=None):
        if self._enabled:
            self._cmd()

    def _on_enter(self, _e=None):
        if self._enabled:
            self.config(relief="raised")

    def _on_leave(self, _e=None):
        if self._enabled:
            self.config(relief="groove")

    def activate(self):
        self._enabled = True
        self.config(bg=self._active_bg, fg=self._active_fg,
                    cursor="arrow", relief="groove")

    def deactivate(self):
        self._enabled = False
        self.config(bg=BTN_GREY_BG, fg=BTN_GREY_FG,
                    cursor="arrow", relief="groove")


# ---------------------------------------------------------------------------
# Helper: rounded-looking card frame
# ---------------------------------------------------------------------------
def Card(parent, **kw) -> tk.Frame:
    kw.setdefault("bg", BG_SURFACE)
    kw.setdefault("highlightbackground", BD_SUBTLE)
    kw.setdefault("highlightthickness", 1)
    kw.setdefault("relief", "flat")
    return tk.Frame(parent, **kw)


def SectionLabel(parent, text: str) -> None:
    """Render a section divider with left label and full-width rule."""
    f = tk.Frame(parent, bg=BG_BASE)
    f.pack(fill="x", padx=20, pady=(16, 6))
    tk.Label(f, text=text, bg=BG_BASE, fg=TX_SECONDARY,
             font=("Helvetica", 10, "bold")).pack(side="left")
    tk.Frame(f, bg=BD_SUBTLE, height=1).pack(
        side="left", fill="x", expand=True, padx=(10, 0), pady=(1, 0))


# ---------------------------------------------------------------------------
# AgentCard
# ---------------------------------------------------------------------------
class AgentCard(tk.Frame):
    def __init__(self, parent, name: str, desc: str,
                 icon: str, color: str, **kw):
        super().__init__(parent, bg=BG_SURFACE,
                         highlightbackground=BD_SUBTLE,
                         highlightthickness=1, **kw)
        self._color = color
        self._spin = 0
        self._scan = 0
        self._animating = False
        self._findings = 0

        # Top accent bar
        tk.Frame(self, bg=color, height=3).pack(fill="x")

        # Header row
        hdr = tk.Frame(self, bg=BG_SURFACE)
        hdr.pack(fill="x", padx=14, pady=(10, 0))
        tk.Label(hdr, text=icon, bg=BG_SURFACE, fg=color,
                 font=("Helvetica", 20)).pack(side="left")
        tk.Label(hdr, text=f"  {name}", bg=BG_SURFACE,
                 fg=CLR_WHITE, font=FONT_H1).pack(side="left")
        self._status = tk.Label(hdr, text="● IDLE", bg=BG_SURFACE,
                                fg=TX_MUTED, font=FONT_SMALL)
        self._status.pack(side="right")

        # Description
        tk.Label(self, text=desc, bg=BG_SURFACE, fg=TX_SECONDARY,
                 font=FONT_SMALL, justify="left",
                 wraplength=200).pack(anchor="w", padx=14, pady=(6, 0))

        # Divider
        tk.Frame(self, bg=BD_SUBTLE, height=1).pack(
            fill="x", padx=14, pady=(8, 0))

        # Animation row
        anim = tk.Frame(self, bg=BG_SURFACE)
        anim.pack(fill="x", padx=14, pady=(6, 0))
        self._spin_lbl = tk.Label(anim, text="", bg=BG_SURFACE,
                                  fg=color, font=FONT_MONO)
        self._spin_lbl.pack(side="left")
        self._scan_lbl = tk.Label(anim, text="", bg=BG_SURFACE,
                                  fg=color, font=FONT_MONO)
        self._scan_lbl.pack(side="left", padx=(6, 0))

        # Findings badge
        self._findings_lbl = tk.Label(self, text="0 findings",
                                      bg=BG_SURFACE, fg=TX_MUTED,
                                      font=FONT_SMALL)
        self._findings_lbl.pack(anchor="w", padx=14, pady=(4, 12))

    def set_status(self, label: str, color: str) -> None:
        self._status.config(text=f"● {label}", fg=color)

    def start_animation(self) -> None:
        self._animating = True
        self._tick()

    def stop_animation(self) -> None:
        self._animating = False
        self._spin_lbl.config(text="")
        self._scan_lbl.config(text="")

    def _tick(self) -> None:
        if not self._animating:
            return
        self._spin_lbl.config(text=SPINNER[self._spin % len(SPINNER)])
        self._scan_lbl.config(text=SCAN[self._scan % len(SCAN)])
        self._spin += 1
        self._scan += 1
        self.after(110, self._tick)

    def add_finding(self) -> None:
        self._findings += 1
        self._findings_lbl.config(
            text=f"{self._findings} finding{'s' if self._findings != 1 else ''}",
            fg=CLR_AMBER)

    def reset(self) -> None:
        self._findings = 0
        self._findings_lbl.config(text="0 findings", fg=TX_MUTED)


# ---------------------------------------------------------------------------
# ConnectorCard
# ---------------------------------------------------------------------------
class ConnectorCard(tk.Frame):
    def __init__(self, parent, name: str, detail: str,
                 icon: str, enabled: bool, **kw):
        super().__init__(parent, bg=BG_SURFACE,
                         highlightbackground=BD_SUBTLE,
                         highlightthickness=1, **kw)
        self._spin = 0
        self._animating = False

        # Left accent strip
        strip_color = CLR_CYAN if enabled else TX_MUTED
        tk.Frame(self, bg=strip_color, width=3).pack(side="left", fill="y")

        body = tk.Frame(self, bg=BG_SURFACE)
        body.pack(side="left", fill="both", expand=True, padx=12, pady=10)

        top = tk.Frame(body, bg=BG_SURFACE)
        top.pack(fill="x")
        tk.Label(top, text=f"{icon}  {name}", bg=BG_SURFACE,
                 fg=CLR_WHITE, font=FONT_H2).pack(side="left")
        color = CLR_GREEN if enabled else TX_MUTED
        label = "ENABLED" if enabled else "DISABLED"
        self._status = tk.Label(top, text=f"● {label}", bg=BG_SURFACE,
                                fg=color, font=FONT_SMALL)
        self._status.pack(side="right")

        tk.Label(body, text=detail, bg=BG_SURFACE,
                 fg=TX_SECONDARY, font=FONT_SMALL).pack(anchor="w", pady=(3, 0))

        self._spin_lbl = tk.Label(body, text="", bg=BG_SURFACE,
                                  fg=CLR_CYAN, font=FONT_MONO)
        self._spin_lbl.pack(anchor="w", pady=(4, 0))

    def set_running(self, running: bool) -> None:
        if running:
            self._status.config(text="● RUNNING", fg=CLR_GREEN)
            self.start_animation()
        else:
            self._status.config(text="● STOPPED", fg=CLR_RED)
            self.stop_animation()

    def start_animation(self) -> None:
        self._animating = True
        self._tick()

    def stop_animation(self) -> None:
        self._animating = False
        self._spin_lbl.config(text="")

    def _tick(self) -> None:
        if not self._animating:
            return
        self._spin_lbl.config(
            text=f"{SPINNER[self._spin % len(SPINNER)]}  "
                 f"{SCAN[self._spin % len(SCAN)]}")
        self._spin += 1
        self.after(110, self._tick)


# ---------------------------------------------------------------------------
# BootPanel
# ---------------------------------------------------------------------------
class BootPanel(tk.Frame):
    STEPS = ["Config","Audit Log","LLM Gateway","Knowledge Base",
             "Authorization","Guardrail Agent","Blue Team Agent",
             "Red Team Agent","DevSecOps Agent","Orchestrator"]

    def __init__(self, parent, **kw):
        super().__init__(parent, bg=BG_BASE, **kw)
        self._spin = 0
        self._running = True

        outer = tk.Frame(self, bg=BG_SURFACE,
                         highlightbackground=BD_ACCENT,
                         highlightthickness=1)
        outer.pack(padx=40, pady=30)

        tk.Label(outer, text="ACDP", bg=BG_SURFACE,
                 fg=CLR_BLUE, font=("Helvetica", 28, "bold")).pack(pady=(24, 2))
        tk.Label(outer, text="Initialising platform…",
                 bg=BG_SURFACE, fg=TX_SECONDARY,
                 font=FONT_BODY).pack(pady=(0, 16))

        steps_frame = tk.Frame(outer, bg=BG_SURFACE)
        steps_frame.pack(padx=30, pady=(0, 10))
        self._icons: dict[str, tk.Label] = {}
        for step in self.STEPS:
            row = tk.Frame(steps_frame, bg=BG_SURFACE)
            row.pack(fill="x", pady=2)
            icon = tk.Label(row, text="○", bg=BG_SURFACE,
                            fg=TX_MUTED, font=FONT_MONO, width=2)
            icon.pack(side="left")
            tk.Label(row, text=f"  {step}", bg=BG_SURFACE,
                     fg=TX_SECONDARY, font=FONT_BODY).pack(side="left")
            self._icons[step] = icon

        self._spin_lbl = tk.Label(outer, text="", bg=BG_SURFACE,
                                  fg=CLR_BLUE, font=FONT_MONO)
        self._spin_lbl.pack(pady=(8, 20))
        self._err = tk.Label(outer, text="", bg=BG_SURFACE,
                             fg=CLR_RED, font=FONT_SMALL, wraplength=400)
        self._err.pack(pady=(0, 14))
        self._tick()

    def mark_ok(self, step: str) -> None:
        if step in self._icons:
            self._icons[step].config(text="✔", fg=CLR_GREEN)

    def mark_err(self, step: str) -> None:
        if step in self._icons:
            self._icons[step].config(text="✘", fg=CLR_RED)

    def show_error(self, msg: str) -> None:
        self._running = False
        self._spin_lbl.config(text="")
        self._err.config(text=f"Error: {msg}")

    def stop(self) -> None:
        self._running = False
        self._spin_lbl.config(text="")

    def _tick(self) -> None:
        if not self._running:
            return
        self._spin_lbl.config(
            text=f"{SPINNER[self._spin % len(SPINNER)]}  booting…")
        self._spin += 1
        self.after(100, self._tick)


# ---------------------------------------------------------------------------
# IngestPanel
# ---------------------------------------------------------------------------
class IngestPanel(tk.Frame):
    CATEGORIES = ["owasp_genai","mitre_attack","mitre_atlas",
                  "playbook","compliance","topology","other"]

    def __init__(self, parent, on_ingest: Callable, **kw):
        super().__init__(parent, bg=BG_SURFACE,
                         highlightbackground=BD_SUBTLE,
                         highlightthickness=1, **kw)
        self._on_ingest = on_ingest
        self._file_path = ""
        self._category_counts: dict[str, int] = {}

        tk.Frame(self, bg=CLR_BLUE, height=3).pack(fill="x")

        body = tk.Frame(self, bg=BG_SURFACE)
        body.pack(fill="x", padx=16, pady=12)

        tk.Label(body, text="Knowledge Ingest", bg=BG_SURFACE,
                 fg=CLR_WHITE, font=FONT_H1).pack(anchor="w")
        tk.Label(body, text="Upload security documents into the RAG knowledge base",
                 bg=BG_SURFACE, fg=TX_SECONDARY, font=FONT_SMALL).pack(
            anchor="w", pady=(2, 10))

        # File row
        file_row = tk.Frame(body, bg=BG_SURFACE)
        file_row.pack(fill="x", pady=(0, 6))

        self._file_var = tk.StringVar(value="No file selected")
        file_bg = tk.Frame(file_row, bg=BG_INPUT,
                           highlightbackground=BD_SUBTLE,
                           highlightthickness=1)
        file_bg.pack(side="left", fill="x", expand=True)
        tk.Label(file_bg, textvariable=self._file_var, bg=BG_INPUT,
                 fg=TX_SECONDARY, font=FONT_SMALL,
                 anchor="w").pack(side="left", padx=8, pady=5, fill="x", expand=True)

        StyledButton(file_row, text="  Browse…  ",
                     command=self._browse,
                     active_bg=BTN_SLATE_BG,
                     active_fg="#e2e8f0",
                     font=FONT_BODY).pack(side="left", padx=(8, 0))

        # Meta row
        meta = tk.Frame(body, bg=BG_SURFACE)
        meta.pack(fill="x", pady=(0, 8))

        tk.Label(meta, text="Source ID", bg=BG_SURFACE,
                 fg=TX_SECONDARY, font=FONT_SMALL).pack(side="left")
        self._src_entry = tk.Entry(meta, bg=BG_INPUT, fg=TX_PRIMARY,
                                   insertbackground=TX_PRIMARY,
                                   relief="flat", font=FONT_BODY,
                                   highlightbackground=BD_SUBTLE,
                                   highlightthickness=1, width=20)
        self._src_entry.pack(side="left", padx=(6, 20), ipady=4)

        tk.Label(meta, text="Category", bg=BG_SURFACE,
                 fg=TX_SECONDARY, font=FONT_SMALL).pack(side="left")
        self._cat = ttk.Combobox(meta, values=self.CATEGORIES,
                                 state="readonly", width=15, font=FONT_BODY)
        self._cat.set("owasp_genai")
        self._cat.pack(side="left", padx=(6, 0))
        # Refresh source ID suggestion whenever category changes
        self._cat.bind("<<ComboboxSelected>>", self._on_category_change)

        # Button row
        btn_row = tk.Frame(body, bg=BG_SURFACE)
        btn_row.pack(fill="x")
        StyledButton(btn_row, text="  ⬆  Ingest  ",
                     command=self._trigger,
                     active_bg=BTN_BLUE_BG,
                     font=FONT_H2).pack(side="left")
        self._status_lbl = tk.Label(btn_row, text="", bg=BG_SURFACE,
                                    fg=CLR_GREEN, font=FONT_SMALL)
        self._status_lbl.pack(side="left", padx=(14, 0))

    # Category → readable base name for source ID
    _CATEGORY_BASES = {
        "owasp_genai":   "owasp-genai",
        "mitre_attack":  "mitre-attack",
        "mitre_atlas":   "mitre-atlas",
        "playbook":      "playbook",
        "compliance":    "compliance",
        "topology":      "topology",
        "other":         "source",
    }

    @staticmethod
    def _auto_source_id(filename: str) -> str:
        """Derive a clean, readable source ID from a filename.

        Examples:
            owasp-genai-top10-2025.pdf  →  owasp-genai-top10-2025
            Enterprise Attack v14.xlsx  →  enterprise-attack-v14
            incident_response.md        →  incident-response
        """
        import re
        stem = Path(filename).stem          # strip extension
        slug = stem.lower()
        slug = re.sub(r"[\s_]+", "-", slug) # spaces/underscores → hyphens
        slug = re.sub(r"[^\w\-]", "", slug) # remove non-alphanumeric except hyphens
        slug = re.sub(r"-{2,}", "-", slug)  # collapse multiple hyphens
        slug = slug.strip("-")
        return slug or "source"

    def _suggest_source_id(self) -> str:
        """Generate source ID from current category + auto-incrementing version."""
        category = self._cat.get()
        base = self._CATEGORY_BASES.get(category, category.replace("_", "-"))
        # Count how many times this category has been used and bump version
        count = self._category_counts.get(category, 0) + 1
        return f"{base}-v{count}"

    def _browse(self) -> None:
        p = filedialog.askopenfilename(
            title="Select knowledge source",
            filetypes=[("Supported","*.pdf *.xlsx *.xls *.txt *.md *.csv"),
                       ("PDF","*.pdf"),("Excel","*.xlsx *.xls"),
                       ("Text / Markdown","*.txt *.md"),("All","*.*")])
        if p:
            self._file_path = p
            self._file_var.set(Path(p).name)
            # Auto-fill source ID from category — user can still edit
            self._src_entry.delete(0, "end")
            self._src_entry.insert(0, self._suggest_source_id())
            self._status_lbl.config(text="")

    def _on_category_change(self, _e=None) -> None:
        """Refresh source ID suggestion when category dropdown changes."""
        self._src_entry.delete(0, "end")
        self._src_entry.insert(0, self._suggest_source_id())

    def _trigger(self) -> None:
        if not self._file_path:
            messagebox.showwarning("No file", "Please select a file first.")
            return
        src = self._src_entry.get().strip()
        if not src:
            messagebox.showwarning("No source ID", "Please enter a source ID.")
            return
        self._status_lbl.config(text="Ingesting…", fg=CLR_AMBER)
        self._on_ingest(self._file_path, src, self._cat.get(), self._done)

    def _done(self, ok: bool, msg: str) -> None:
        if ok:
            # Bump version counter so next ingest of same category gets v2, v3…
            cat = self._cat.get()
            self._category_counts[cat] = self._category_counts.get(cat, 0) + 1
            # Pre-fill the next suggested ID ready for another upload
            self._src_entry.delete(0, "end")
            self._src_entry.insert(0, self._suggest_source_id())
        self._status_lbl.config(
            text=("✔  " if ok else "✘  ") + msg,
            fg=CLR_GREEN if ok else CLR_RED)


# ---------------------------------------------------------------------------
# AuditTail  (resizable via sash handle)
# ---------------------------------------------------------------------------
class AuditTail(tk.Frame):
    def __init__(self, parent, **kw):
        super().__init__(parent, bg=BG_SURFACE,
                         highlightbackground=BD_SUBTLE,
                         highlightthickness=1, **kw)
        self._n = 0

        # Header
        hdr = tk.Frame(self, bg=BG_RAISED)
        hdr.pack(fill="x")
        tk.Frame(hdr, bg=CLR_GREEN, width=3).pack(side="left", fill="y")
        tk.Label(hdr, text="  Live Audit Log", bg=BG_RAISED,
                 fg=CLR_WHITE, font=FONT_H2).pack(side="left", padx=(4, 0), pady=8)
        self._count = tk.Label(hdr, text="0 events", bg=BG_RAISED,
                               fg=TX_SECONDARY, font=FONT_SMALL)
        self._count.pack(side="right", padx=12)

        # Text area
        text_frame = tk.Frame(self, bg=BG_BASE)
        text_frame.pack(fill="both", expand=True)

        self._text = tk.Text(text_frame, bg=BG_BASE, fg=TX_PRIMARY,
                             font=FONT_MONO, relief="flat",
                             state="disabled", wrap="none",
                             selectbackground=BD_ACCENT,
                             insertbackground=TX_PRIMARY)
        xsb = tk.Scrollbar(text_frame, orient="horizontal",
                            command=self._text.xview, bg=BG_SURFACE)
        ysb = tk.Scrollbar(text_frame, orient="vertical",
                            command=self._text.yview, bg=BG_SURFACE)
        self._text.configure(xscrollcommand=xsb.set,
                             yscrollcommand=ysb.set)
        ysb.pack(side="right", fill="y")
        xsb.pack(side="bottom", fill="x")
        self._text.pack(side="left", fill="both", expand=True)

        # Colour tags
        self._text.tag_config("info",    foreground=TX_SECONDARY)
        self._text.tag_config("success", foreground=CLR_GREEN)
        self._text.tag_config("warning", foreground=CLR_AMBER)
        self._text.tag_config("error",   foreground=CLR_RED)
        self._text.tag_config("finding", foreground="#c084fc")

        # Resize handle (drag up/down to change height)
        handle = tk.Frame(self, bg=BD_ACCENT, cursor="sb_v_double_arrow", height=6)
        handle.pack(fill="x", side="bottom")
        handle.bind("<ButtonPress-1>",   self._resize_start)
        handle.bind("<B1-Motion>",       self._resize_drag)

        self._resize_y = 0
        self._resize_h = 0

    def append(self, line: str, level: str = "info") -> None:
        self._text.configure(state="normal")
        ts = datetime.datetime.now().strftime("%H:%M:%S")
        self._text.insert("end", f"[{ts}]  {line}\n", level)
        self._text.see("end")
        self._text.configure(state="disabled")
        self._n += 1
        self._count.config(text=f"{self._n} events")

    def _resize_start(self, e) -> None:
        self._resize_y = e.y_root
        self._resize_h = self.winfo_height()

    def _resize_drag(self, e) -> None:
        delta = e.y_root - self._resize_y
        new_h = max(80, self._resize_h + delta)
        self.configure(height=new_h)
        self.pack_propagate(False)


# ---------------------------------------------------------------------------
# Dashboard — main window
# ---------------------------------------------------------------------------
class Dashboard:
    def __init__(self, config_path: str):
        self._config_path = config_path
        self._platform: Any = None
        self._running = False
        self._q: queue.Queue = queue.Queue()

        self.root = tk.Tk()
        self.root.title("ACDP — Autonomous Cyber Defense Platform")
        self.root.configure(bg=BG_BASE)
        self.root.geometry("1160x900")
        self.root.minsize(960, 720)
        self._apply_style()
        self._build_header()
        self._build_boot_screen()
        self.root.after(100, self._poll)
        threading.Thread(target=self._boot_platform, daemon=True).start()

    # ------------------------------------------------------------------
    def _apply_style(self) -> None:
        s = ttk.Style(self.root)
        s.theme_use("clam")
        s.configure("TCombobox",
                    fieldbackground=BG_INPUT, background=BG_INPUT,
                    foreground=TX_PRIMARY, selectbackground=BG_INPUT,
                    selectforeground=TX_PRIMARY, borderwidth=0,
                    arrowcolor=TX_SECONDARY)
        s.map("TCombobox",
              fieldbackground=[("readonly", BG_INPUT)],
              foreground=[("readonly", TX_PRIMARY)])

    # ------------------------------------------------------------------
    # HEADER — centered hero + version top-right
    # ------------------------------------------------------------------
    def _build_header(self) -> None:
        bar = tk.Frame(self.root, bg=BG_HEADER)
        bar.pack(fill="x", side="top")

        # Version badge — top right
        tk.Label(bar, text="v0.1.0", bg=BG_HEADER,
                 fg=TX_MUTED, font=FONT_VER).pack(
            side="right", padx=18, pady=14)

        # Config path — top right (secondary)
        tk.Label(bar, text=f"config: {self._config_path}",
                 bg=BG_HEADER, fg=TX_MUTED,
                 font=FONT_VER).pack(side="right", padx=(0, 4), pady=14)

        # Center column for hero title
        center = tk.Frame(bar, bg=BG_HEADER)
        center.pack(expand=True)

        tk.Label(center, text="ACDP", bg=BG_HEADER,
                 fg=CLR_BLUE, font=FONT_HERO).pack()
        tk.Label(center, text="Autonomous Cyber Defense Platform",
                 bg=BG_HEADER, fg=TX_SECONDARY,
                 font=FONT_SUB).pack(pady=(2, 14))

        # Thin accent rule under header
        tk.Frame(self.root, bg=CLR_BLUE, height=2).pack(fill="x")

    # ------------------------------------------------------------------
    # Boot screen
    # ------------------------------------------------------------------
    def _build_boot_screen(self) -> None:
        self._boot_frame = tk.Frame(self.root, bg=BG_BASE)
        self._boot_frame.pack(fill="both", expand=True)
        self._boot_panel = BootPanel(self._boot_frame)
        self._boot_panel.place(relx=0.5, rely=0.45, anchor="center")

    # ------------------------------------------------------------------
    # Platform boot (background)
    # ------------------------------------------------------------------
    def _boot_platform(self) -> None:
        try:
            from acdp.config import ConfigLoader
            from acdp.main import Platform
            self._q.put(("boot_ok", "Config"))
            config = ConfigLoader().load(Path(self._config_path))
            self._q.put(("boot_ok", "Audit Log"))
            platform = Platform.from_config(config)
            for step in BootPanel.STEPS[2:]:
                self._q.put(("boot_ok", step))
                time.sleep(0.04)
            self._platform = platform
            self._q.put(("boot_done", None))
        except Exception as exc:
            self._q.put(("boot_error", str(exc)))

    # ------------------------------------------------------------------
    # Main dashboard (built after boot)
    # ------------------------------------------------------------------
    def _build_main(self) -> None:
        self._boot_panel.stop()
        self._boot_frame.destroy()

        cfg = self._platform.config

        # Outer paned window so audit log is vertically resizable
        paned = tk.PanedWindow(self.root, orient="vertical",
                               bg=BG_BASE, sashwidth=8,
                               sashrelief="flat",
                               sashpad=2,
                               handlesize=0)
        paned.pack(fill="both", expand=True)

        # Top scrollable content pane
        top_pane = tk.Frame(paned, bg=BG_BASE)
        paned.add(top_pane, stretch="always", minsize=400)

        canvas = tk.Canvas(top_pane, bg=BG_BASE, highlightthickness=0)
        vsb = tk.Scrollbar(top_pane, orient="vertical",
                           command=canvas.yview, bg=BG_SURFACE)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        inner = tk.Frame(canvas, bg=BG_BASE)
        win = canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>",
                   lambda e: canvas.configure(
                       scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>",
                    lambda e: canvas.itemconfig(win, width=e.width))
        canvas.bind_all("<MouseWheel>",
                        lambda e: canvas.yview_scroll(
                            int(-1*(e.delta/120)), "units"))

        self._populate_content(inner, cfg)

        # Bottom resizable audit log pane
        log_pane = tk.Frame(paned, bg=BG_BASE)
        paned.add(log_pane, stretch="never", minsize=120)

        self._log_tail = AuditTail(log_pane)
        self._log_tail.pack(fill="both", expand=True,
                            padx=20, pady=(8, 12))

        # Set initial sash position (audit log gets ~220px)
        self.root.update_idletasks()
        total = paned.winfo_height()
        paned.sash_place(0, 0, max(total - 220, 400))

        self._log("success", "Platform booted successfully")
        self._log("info",
                  "Enable connectors in config.yaml, then click Start All Connectors")

    def _populate_content(self, parent: tk.Frame, cfg) -> None:
        # ── Agents ──────────────────────────────────────────────────
        SectionLabel(parent, "AGENTS")
        grid = tk.Frame(parent, bg=BG_BASE)
        grid.pack(fill="x", padx=20, pady=(0, 4))
        grid.columnconfigure((0,1,2,3), weight=1, uniform="ag")

        agents = [
            ("Guardrail", "Inline LLM firewall\nBlocks prompt injection\nScrubs PII & secrets",  "🛡"),
            ("Blue Team", "Telemetry parsing\nAnomaly detection\nContainment requests",           "🔵"),
            ("Red Team",  "Authorized probing\nAdversary simulation\nWeakness discovery",         "🔴"),
            ("DevSecOps", "Code vulnerability\nPatch generation\nPR-based remediation",           "🔧"),
        ]
        self._agent_cards: dict[str, AgentCard] = {}
        for col, (name, desc, icon) in enumerate(agents):
            c = AgentCard(grid, name, desc, icon=icon,
                          color=AGENT_COLORS[name])
            c.grid(row=0, column=col, sticky="nsew", padx=6, pady=4)
            self._agent_cards[name] = c

        # ── Connectors ───────────────────────────────────────────────
        SectionLabel(parent, "CONNECTORS")
        cgrid = tk.Frame(parent, bg=BG_BASE)
        cgrid.pack(fill="x", padx=20, pady=(0, 4))
        cgrid.columnconfigure((0,1,2), weight=1, uniform="cn")

        connectors_info = [
            ("GuardrailProxy",
             f"HTTP firewall  ·  port {cfg.connectors.guardrail_proxy.port}",
             "🔒", cfg.connectors.guardrail_proxy.enabled),
            ("LogStream",
             f"Log monitor  ·  mode: {cfg.connectors.log_stream.mode}",
             "📡", cfg.connectors.log_stream.enabled),
            ("Scheduler",
             f"Red Team scan  ·  mode: {cfg.connectors.scheduler.mode}",
             "⏱", cfg.connectors.scheduler.enabled),
        ]
        self._conn_cards: dict[str, ConnectorCard] = {}
        for col, (name, detail, icon, enabled) in enumerate(connectors_info):
            c = ConnectorCard(cgrid, name, detail,
                              icon=icon, enabled=enabled)
            c.grid(row=0, column=col, sticky="nsew", padx=6, pady=4)
            self._conn_cards[name] = c

        # ── Control bar ──────────────────────────────────────────────
        ctrl = tk.Frame(parent, bg=BG_BASE)
        ctrl.pack(fill="x", padx=20, pady=(8, 4))

        self._start_btn = StyledButton(
            ctrl, text="  ▶  Start All Connectors  ",
            command=self._start_all,
            active_bg=BTN_START_BG,
            font=FONT_H2)
        self._start_btn.pack(side="left")

        self._stop_btn = StyledButton(
            ctrl, text="  ■  Stop All  ",
            command=self._stop_all,
            active_bg=BTN_STOP_BG,
            font=FONT_H2)
        self._stop_btn.pack(side="left", padx=(12, 0))
        self._stop_btn.deactivate()   # greyed out initially

        self._plat_status = tk.Label(ctrl, text="● Platform ready",
                                     bg=BG_BASE, fg=CLR_GREEN,
                                     font=FONT_H2)
        self._plat_status.pack(side="right")

        # ── Ingest + spacer (left half only) ─────────────────────────
        SectionLabel(parent, "KNOWLEDGE INGEST")
        ingest_row = tk.Frame(parent, bg=BG_BASE)
        ingest_row.pack(fill="x", padx=20, pady=(0, 8))
        ingest_row.columnconfigure(0, weight=1)
        ingest_row.columnconfigure(1, weight=1)

        ingest = IngestPanel(ingest_row, on_ingest=self._do_ingest)
        ingest.grid(row=0, column=0, sticky="nsew", padx=(0, 8))

        # Right half: placeholder / tips card
        tips = tk.Frame(ingest_row, bg=BG_SURFACE,
                        highlightbackground=BD_SUBTLE,
                        highlightthickness=1)
        tips.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        tk.Frame(tips, bg=BD_SUBTLE, height=3).pack(fill="x")
        tk.Label(tips, text="Ingest order", bg=BG_SURFACE,
                 fg=CLR_WHITE, font=FONT_H2).pack(
            anchor="w", padx=16, pady=(12, 4))
        tips_text = (
            "1.  owasp_genai   — trains Guardrail to block\n"
            "     prompt injection & jailbreaks\n\n"
            "2.  mitre_attack  — ATT&CK adversary knowledge\n"
            "     for Red Team probe context\n\n"
            "3.  playbook      — incident runbooks for Blue\n"
            "     Team containment provenance\n\n"
            "4.  compliance    — secure-coding standards for\n"
            "     DevSecOps PR descriptions"
        )
        tk.Label(tips, text=tips_text, bg=BG_SURFACE,
                 fg=TX_SECONDARY, font=FONT_SMALL,
                 justify="left", anchor="nw").pack(
            anchor="w", padx=16, pady=(0, 14))


    # ------------------------------------------------------------------
    # Start / Stop
    # ------------------------------------------------------------------
    def _start_all(self) -> None:
        if self._running:
            return
        self._running = True
        self._start_btn.deactivate()
        self._stop_btn.activate()
        self._plat_status.config(text="● Starting…", fg=CLR_AMBER)
        for card in self._agent_cards.values():
            card.set_status("ACTIVE", CLR_GREEN)
            card.start_animation()
        cfg = self._platform.config.connectors
        enabled_map = {"GuardrailProxy": cfg.guardrail_proxy.enabled,
                       "LogStream":      cfg.log_stream.enabled,
                       "Scheduler":      cfg.scheduler.enabled}
        for name, card in self._conn_cards.items():
            if enabled_map.get(name):
                card.set_running(True)
        self._log("info", "Starting all enabled connectors…")
        threading.Thread(target=self._start_thread, daemon=True).start()

    def _start_thread(self) -> None:
        try:
            self._platform.start_connectors()
            self._q.put(("status", "● All connectors running", CLR_GREEN))
            self._q.put(("log", "success", "All connectors started"))
        except Exception as exc:
            self._q.put(("status", f"● Error: {exc}", CLR_RED))
            self._q.put(("log", "error", f"Connector error: {exc}"))

    def _stop_all(self) -> None:
        if not self._running:
            return
        self._running = False
        self._stop_btn.deactivate()
        self._start_btn.activate()
        self._plat_status.config(text="● Stopping…", fg=CLR_AMBER)
        self._log("warning", "Stopping all connectors…")
        for card in self._agent_cards.values():
            card.set_status("IDLE", TX_MUTED)
            card.stop_animation()
        for card in self._conn_cards.values():
            card.set_running(False)
        threading.Thread(target=self._stop_thread, daemon=True).start()

    def _stop_thread(self) -> None:
        try:
            self._platform.stop_connectors()
            self._q.put(("status", "● Platform ready", CLR_GREEN))
            self._q.put(("log", "info", "All connectors stopped cleanly"))
        except Exception as exc:
            self._q.put(("log", "error", f"Stop error: {exc}"))
        self._q.put(("reset_buttons", None, None))

    # ------------------------------------------------------------------
    # Ingest
    # ------------------------------------------------------------------
    def _do_ingest(self, file_path: str, source_id: str,
                   category: str, cb: Callable) -> None:
        def _work():
            try:
                from acdp.models import KnowledgeSource, SourceCategory
                content = _read_file(Path(file_path))
                self._platform.ingestion_pipeline.ingest(
                    KnowledgeSource(source_id=source_id,
                                    category=SourceCategory(category),
                                    content=content))
                name = Path(file_path).name
                self._q.put(("log", "success",
                              f"Ingested {name}  [{category}]  id={source_id}"))
                cb(True, f"Ingested {name}")
            except Exception as exc:
                self._q.put(("log", "error", f"Ingest error: {exc}"))
                cb(False, str(exc))
        self._log("info",
                  f"Ingesting {Path(file_path).name} as [{category}]…")
        threading.Thread(target=_work, daemon=True).start()

    # ------------------------------------------------------------------
    # Queue polling
    # ------------------------------------------------------------------
    def _poll(self) -> None:
        try:
            while True:
                msg = self._q.get_nowait()
                kind = msg[0]
                if kind == "boot_ok":
                    self._boot_panel.mark_ok(msg[1])
                elif kind == "boot_done":
                    self._build_main()
                elif kind == "boot_error":
                    self._boot_panel.show_error(msg[1])
                elif kind == "log":
                    self._log(msg[1], msg[2])
                elif kind == "status":
                    self._plat_status.config(text=msg[1], fg=msg[2])
                elif kind == "reset_buttons":
                    self._start_btn.activate()
                    self._stop_btn.deactivate()
        except queue.Empty:
            pass
        self.root.after(100, self._poll)

    def _log(self, level: str, text: str) -> None:
        try:
            self._log_tail.append(text, level)
        except AttributeError:
            pass

    def run(self) -> None:
        self.root.mainloop()


# ---------------------------------------------------------------------------
# File reader helper
# ---------------------------------------------------------------------------
def _read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import pypdf
            r = pypdf.PdfReader(str(path))
            return "\n".join(p.extract_text() or "" for p in r.pages)
        except ImportError:
            raise RuntimeError("pypdf not installed — run: pip install pypdf")
    if suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join("" if c is None else str(c) for c in row)
                    if line.strip():
                        lines.append(line)
            return "\n".join(lines)
        except ImportError:
            raise RuntimeError(
                "openpyxl not installed — run: pip install openpyxl")
    return path.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
def launch_dashboard(config_path: str) -> None:
    """Open the GUI dashboard and block until the window is closed."""
    Dashboard(config_path).run()
